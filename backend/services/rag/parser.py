"""
UPLOAD -> FILE VALIDATION -> TEXT EXTRACTION -> STRUCTURE DETECTION
(spec section 6). Returns a list of "pages" (or slide/section units) each
carrying whatever structural metadata the format can offer, so the chunker
can attach page/chapter/heading info per chunk.
"""
from __future__ import annotations
import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    heading: str | None = None


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage] = field(default_factory=list)
    page_count: int = 0
    ocr_used: bool = False  # True if any page's text came from OCR rather than the embedded text layer
    ocr_available: bool = True  # False if this environment has no tesseract/poppler installed


class ParsingError(Exception):
    pass


HEADING_RE = re.compile(r"^\s{0,3}(chapter|section|part)\s+\d+[:.\s]", re.IGNORECASE)


def _detect_heading(line: str) -> str | None:
    line = line.strip()
    if not line:
        return None
    if HEADING_RE.match(line):
        return line
    # Short, title-cased, no trailing period -> likely a heading
    if len(line) < 80 and not line.endswith(".") and line[:1].isupper() and line.count(" ") < 10:
        return line
    return None


# A page with fewer than this many extracted characters is treated as
# "likely scanned" and sent through OCR instead (a page number or stray
# header can produce a handful of characters even on a pure-image page).
OCR_TRIGGER_CHAR_THRESHOLD = 20


def _ocr_available() -> bool:
    import shutil
    return shutil.which("tesseract") is not None and shutil.which("pdftoppm") is not None


def _ocr_pdf_page(path: Path, page_number: int) -> str:
    """Rasterizes one PDF page and runs Tesseract OCR on it. Returns '' on
    any failure rather than raising — OCR is a best-effort enhancement, and
    a page that can't be OCR'd shouldn't kill ingestion of the whole document."""
    try:
        from pdf2image import convert_from_path
        import pytesseract

        images = convert_from_path(str(path), first_page=page_number, last_page=page_number, dpi=200)
        if not images:
            return ""
        return pytesseract.image_to_string(images[0]) or ""
    except Exception:
        return ""


def parse_pdf(path: Path) -> ExtractedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as e:
        raise ParsingError(f"Could not open PDF: {e}") from e

    ocr_available = _ocr_available()
    pages = []
    current_heading = None
    any_ocr_used = False

    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            raise ParsingError(f"Failed to extract text from page {i}: {e}") from e

        if len(text.strip()) < OCR_TRIGGER_CHAR_THRESHOLD and ocr_available:
            ocr_text = _ocr_pdf_page(path, i)
            if len(ocr_text.strip()) > len(text.strip()):
                text = ocr_text
                any_ocr_used = True

        first_line = text.strip().split("\n")[0] if text.strip() else ""
        h = _detect_heading(first_line)
        if h:
            current_heading = h
        pages.append(ExtractedPage(page_number=i, text=text, heading=current_heading))

    doc = ExtractedDocument(pages=pages, page_count=len(pages))
    doc.ocr_used = any_ocr_used
    doc.ocr_available = ocr_available
    return doc


def parse_docx(path: Path) -> ExtractedDocument:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise ParsingError(f"Could not open DOCX: {e}") from e

    pages = []
    buf: list[str] = []
    current_heading = None
    page_num = 1
    for para in d.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            if buf:
                pages.append(ExtractedPage(page_number=page_num, text="\n".join(buf), heading=current_heading))
                buf = []
                page_num += 1
            current_heading = para.text.strip()
        if para.text.strip():
            buf.append(para.text)
    if buf:
        pages.append(ExtractedPage(page_number=page_num, text="\n".join(buf), heading=current_heading))
    if not pages:
        raise ParsingError("DOCX contained no extractable text.")
    return ExtractedDocument(pages=pages, page_count=len(pages))


def parse_pptx(path: Path) -> ExtractedDocument:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise ParsingError(f"Could not open PPTX: {e}") from e

    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        heading = None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = shape.text_frame.text.strip()
            if not t:
                continue
            if heading is None and shape == slide.shapes.title:
                heading = t
            texts.append(t)
        pages.append(ExtractedPage(page_number=i, text="\n".join(texts), heading=heading or f"Slide {i}"))
    return ExtractedDocument(pages=pages, page_count=len(pages))


def parse_text(path: Path) -> ExtractedDocument:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        raise ParsingError(f"Could not read text file: {e}") from e
    # Split plain text into pseudo-pages of ~2000 chars on paragraph boundaries
    paras = content.split("\n\n")
    pages, buf, size, page_num = [], [], 0, 1
    for p in paras:
        buf.append(p)
        size += len(p)
        if size > 2000:
            pages.append(ExtractedPage(page_number=page_num, text="\n\n".join(buf)))
            buf, size, page_num = [], 0, page_num + 1
    if buf:
        pages.append(ExtractedPage(page_number=page_num, text="\n\n".join(buf)))
    return ExtractedDocument(pages=pages, page_count=len(pages))


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".doc": parse_docx,
    ".pptx": parse_pptx,
    ".ppt": parse_pptx,
    ".txt": parse_text,
    ".md": parse_text,
}


def parse_document(path: Path) -> ExtractedDocument:
    ext = path.suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ParsingError(f"Unsupported file type: {ext}")

    doc = parser(path)
    if not any(p.text.strip() for p in doc.pages):
        if ext == ".pdf" and not doc.ocr_available:
            raise ParsingError(
                "No extractable text found, and OCR is unavailable in this "
                "deployment (tesseract/poppler not installed). This looks like "
                "a scanned/image-only PDF — install tesseract-ocr and "
                "poppler-utils to enable OCR (see backend/Dockerfile)."
            )
        raise ParsingError(
            "No extractable text found, even after attempting OCR. The file "
            "may be blank, corrupted, or contain only non-text content."
            if ext == ".pdf"
            else "No extractable text found. The file may be a scanned/image-only "
                 "document — OCR is only wired for PDFs currently."
        )
    return doc

